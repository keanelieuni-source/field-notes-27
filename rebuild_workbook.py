"""Rebuild history.xlsx from every entry in snapshots.json. pip install openpyxl"""
import base64
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent
data = json.loads((ROOT / "snapshots.json").read_text(encoding="utf-8"))
wb = Workbook()
wb.remove(wb.active)
rows = {name: [] for name in ["Latitude plans", "Latitude locations", "OpenRouter models", "OpenRouter shares", "Ramp adoption", "Ramp financing", "Ramp token spend"]}
for key, s in sorted(data["snapshots"].items()):
    for p in s["latitude"]["plans"]:
        row = {"snapshotDate": key, **{k: v for k, v in p.items() if not isinstance(v, (list, dict))}}
        for field in ["cheapestLocations", "mostExpensiveLocations"]:
            row[field] = ", ".join(p.get(field, []))
        rows["Latitude plans"].append(row)
        for loc in p.get("locations", []):
            rows["Latitude locations"].append({"snapshotDate": key, "plan": p["id"], **loc})
    o, r = s["openrouter"], s["ramp"]
    rows["OpenRouter models"] += [{"snapshotDate": key, "usageThrough": o["usageThrough"], **m} for m in o["weekTopModels"]]
    rows["OpenRouter shares"] += [{"snapshotDate": key, "weekBeginning": o["marketShareWeekOf"], **m} for m in o["marketShare"]]
    rows["Ramp adoption"].append({"snapshotDate": key, "period": r["period"], **r["adoption"], **{k: r.get(k) for k in ["spendPerEmployee", "top1SpendPerEmployee", "top10SpendPerEmployee"]}})
    rows["Ramp financing"] += [{"snapshotDate": key, "period": r["period"], **m} for m in r["financingStatus"]]
    rows["Ramp token spend"] += [{"snapshotDate": key, "periodEnd": r["tokenSpend"]["periodEnd"], "unit": r["tokenSpend"]["unit"], **m} for m in r["tokenSpend"]["makers"]]
rows["Read me"] = [{"note": x} for x in ["USD per whole server. Annual totals remain annual. Missing values are blank, never zero.", "Source dates can differ from snapshot dates. Compare at least two dates before claiming a change.", "Latitude includes out-of-stock offers. VM location availability was not verified.", "Ramp token-spend volume is indexed spend, not raw tokens or USD.", "OpenRouter rankings: CC BY 4.0. https://openrouter.ai/rankings"]]
for name, records in rows.items():
    ws = wb.create_sheet(name)
    if not records:
        continue
    headers = list(dict.fromkeys(k for row in records for k in row))
    ws.append(headers)
    for row in records:
        ws.append([json.dumps(row.get(k)) if isinstance(row.get(k), (list, dict)) else row.get(k) for k in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2B45C4")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(58, max(14, max(len(str(c.value or "")) for c in col) + 2))
wb.save(ROOT / "history.xlsx")
(ROOT / "workbook.b64").write_text(base64.b64encode((ROOT / "history.xlsx").read_bytes()).decode() + "\n", encoding="ascii")
print("Rebuilt history.xlsx from", len(data["snapshots"]), "snapshot(s)")
